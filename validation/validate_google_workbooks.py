from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHEETS = ROOT / "google" / "sheets"


def headers(workbook_name: str, sheet_name: str) -> list[str]:
    workbook = load_workbook(SHEETS / workbook_name, read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name]
        return [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    finally:
        workbook.close()


def assert_headers(workbook_name: str, sheet_name: str, expected: list[str]) -> None:
    actual = headers(workbook_name, sheet_name)
    missing = [name for name in expected if name not in actual]
    assert not missing, f"{workbook_name}/{sheet_name} missing headers: {missing}"


def main() -> None:
    assert_headers("air-test.xlsx", "records-Air", ["worksheetCreate", "building"])
    assert_headers("air-test.xlsx", "records-CA Gass", ["worksheetCreate", "building"])
    assert_headers("RPP2-air-record.xlsx", "records_em_B10", ["worksheetNo", "samplesJson"])
    assert_headers("RPP2-air-record.xlsx", "records_ca_OT", ["worksheetNo", "samplesJson"])

    assert_headers("water-r.xlsx", "prw-pw", ["samplingDate", "worksheet No.", "building"])
    assert_headers("water-r.xlsx", "wfi-pus", ["samplingDate", "worksheet  No.", "building"])
    assert_headers("RPP2-water-record.xlsx", "records_pw_prw_B10", ["worksheetNo", "samplesJson"])
    assert_headers("RPP2-water-record.xlsx", "records_wfi_B16", ["worksheetNo", "samplesJson"])

    assert_headers("Testing.xlsx", "CV", [
        "worksheetNo", "worksheetCreate", "Bld", "Sampling date", "Samp-Method",
        "Test-Method", "CV/CEHT", "Normal/ReSamp", "ครั้งที่", "syncStatus",
    ])
    assert_headers("RPP2-cv-record.xlsx", "records_cv_contact_B10", ["worksheetNo", "samplesJson"])
    assert_headers("RPP2-cv-record.xlsx", "record_cv_rinse_B10", ["worksheetNo", "samplesJson"])

    print("Google workbook sheet/header contracts: PASS")


if __name__ == "__main__":
    main()
